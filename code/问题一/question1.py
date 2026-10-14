import math as m

import openpyxl

depth_list = []
l_width_list = []
r_width_list = []
d_values = [-800, -600, -400, -200, 0, 200, 400, 600, 800]


def seek_wd(d1, a, b):
    depth = -d1 * m.tan(a) + 70
    l_width = (depth / m.cos((b / 2) + a)) * m.sin(b / 2)*m.cos(a)
    r_width = (depth / m.cos(a - (b / 2))) * m.sin(b / 2)*m.cos(a)
    return depth, l_width, r_width


l = 200
for d in d_values:
    dep, l_wid, r_wid = seek_wd(d, m.radians(1.5), m.radians(120))
    depth_list.append(dep)
    l_width_list.append(l_wid)
    r_width_list.append(r_wid)
width = []
for i in range(len(l_width_list)):
    width.append(l_width_list[i] + r_width_list[i])
mid_width=[]
for i in range(len(l_width_list)-1):
    mid_width.append(l_width_list[i+1] + r_width_list[i])
result = []
for i in range(len(width)-1):
    result.append((1 - (l / mid_width[i]))*100)
workbook = openpyxl.load_workbook('result1.xlsx')

sheet = workbook.active

start_row = 2
start_column = 2

row_index = 2
for col_index, value in enumerate(depth_list, start=2):
    sheet.cell(row=row_index, column=col_index, value=value)
row_index = 3
for col_index, value in enumerate(width, start=2):
    sheet.cell(row=row_index, column=col_index, value=value)
row_index = 4
for col_index, value in enumerate(result, start=3):
    sheet.cell(row=row_index, column=col_index, value=value)
workbook.save('result1.xlsx')