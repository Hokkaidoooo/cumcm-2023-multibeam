"""带输入参数和历史来源的可复现实验。"""
from hashlib import sha256
from importlib.metadata import version
from math import cos, isclose, isfinite, radians, sin, tan
from pathlib import Path
import platform

from .coverage import audit_strips, parallel_lines
from .geometry import PlanarSeabed, cross_section_swath
from .legacy import question2_value, question3_strips

NAUTICAL_MILE_M = 1852
Q3_WIDTH_M = 4 * NAUTICAL_MILE_M
Q3_LENGTH_M = 2 * NAUTICAL_MILE_M
Q3_OVERLAPS = (.105, .12, .15, .18)


def file_record(path: Path, root: Path) -> dict:
    return {"path": path.relative_to(root).as_posix(), "bytes": path.stat().st_size, "sha256": sha256(path.read_bytes()).hexdigest()}


def environment() -> dict:
    return {"python": platform.python_version(), "platform": platform.platform(), "openpyxl": version("openpyxl")}


def question1() -> dict:
    seabed = PlanarSeabed(70, -tan(radians(1.5)))
    rows = []
    previous = None
    for distance in range(-800, 801, 200):
        swath = cross_section_swath(seabed.depth(distance), seabed.gradient_x)
        rate = None if previous is None else 1 - 200 / (previous.positive_m + swath.negative_m)
        rows.append({"distance_m": distance, "depth_m": swath.depth_m, "horizontal_width_m": swath.horizontal_width_m, "seabed_width_m": swath.seabed_width_m, "paper_overlap_rate": rate})
        previous = swath
    return {"parameters": {"depth_origin_m": 70, "slope_deg": -1.5, "opening_deg": 120, "distance_unit": "horizontal metre"}, "rows": rows}


def question2(root: Path) -> dict:
    import openpyxl

    source = root / "code" / "问题二" / "result2.xlsx"
    book = openpyxl.load_workbook(source, data_only=True, read_only=True)
    sheet = book["Sheet1"]
    seabed = PlanarSeabed(120, tan(radians(1.5)))
    rows = []
    for row, beta in enumerate(range(0, 360, 45), 3):
        for column, step in enumerate(range(8), 3):
            distance = step * .3 * NAUTICAL_MILE_M
            heading = radians(beta)
            x, y = distance * cos(heading), distance * sin(heading)
            swath = seabed.swath(x, y, beta)
            old_projected, old_unprojected = question2_value(beta, distance)
            historical = sheet.cell(row, column).value
            if not isinstance(historical, (int, float)):
                raise ValueError(f"Sheet1!{sheet.cell(row, column).coordinate}: numeric value required")
            rows.append({
                "cell": f"Sheet1!{sheet.cell(row, column).coordinate}",
                "beta_deg": beta, "distance_nm": step * .3, "distance_m": distance,
                "x_m": x, "y_m": y, "depth_m": swath.depth_m, "cross_slope": swath.cross_slope,
                "historical_workbook_m": historical, "legacy_script_horizontal_m": old_projected,
                "legacy_script_without_projection_m": old_unprojected,
                "corrected_horizontal_m": swath.horizontal_width_m, "corrected_seabed_m": swath.seabed_width_m,
                "legacy_minus_workbook_m": old_projected - historical,
                "corrected_horizontal_minus_workbook_m": swath.horizontal_width_m - historical,
            })
    book.close()
    return {
        "source": file_record(source, root), "environment": environment(),
        "parameters": {"depth_origin_m": 120, "slope_deg": 1.5, "opening_deg": 120, "beta_zero": "+x, increasing depth", "distance_unit": "horizontal nautical mile", "nautical_mile_m": NAUTICAL_MILE_M},
        "tolerance": {"relative": 1e-9, "absolute_m": 1e-6},
        "legacy_mismatch_count": sum(not isclose(r["legacy_script_horizontal_m"], r["historical_workbook_m"], rel_tol=1e-9, abs_tol=1e-6) for r in rows),
        "without_projection_match_count": sum(isclose(r["legacy_script_without_projection_m"], r["historical_workbook_m"], rel_tol=1e-9, abs_tol=1e-6) for r in rows),
        "legacy_max_absolute_difference_m": max(abs(r["legacy_minus_workbook_m"]) for r in rows),
        "corrected_horizontal_mismatch_count": sum(not isclose(r["corrected_horizontal_m"], r["historical_workbook_m"], rel_tol=1e-9, abs_tol=1e-6) for r in rows),
        "corrected_max_absolute_difference_m": max(abs(r["corrected_horizontal_minus_workbook_m"]) for r in rows),
        "rows": rows,
    }


def question3() -> dict:
    slope = -tan(radians(1.5))
    seabed = PlanarSeabed(110 - slope * Q3_WIDTH_M / 2, slope)
    old = audit_strips(question3_strips(), (0, Q3_WIDTH_M))
    old["line_count"] = len(old["strips"])
    old["total_length_m"] = old["line_count"] * Q3_LENGTH_M
    candidates = []
    for overlap in Q3_OVERLAPS:
        strips = parallel_lines(seabed, Q3_WIDTH_M, overlap)
        audit = audit_strips(strips, (0, Q3_WIDTH_M))
        audit.update({"facing_overlap": overlap, "line_count": len(strips), "total_length_m": len(strips) * Q3_LENGTH_M})
        candidates.append(audit)
    feasible = [c for c in candidates if c["fully_covered"] and c["overlap_satisfied"]]
    if not feasible:
        raise ValueError("no candidate satisfies coverage and overlap constraints")
    best = min(feasible, key=lambda c: c["total_length_m"])
    return {
        "parameters": {"width_m": Q3_WIDTH_M, "line_length_m": Q3_LENGTH_M, "center_depth_m": 110, "west_depth_m": seabed.depth_origin_m, "east_depth_m": seabed.depth(Q3_WIDTH_M), "slope_deg": -1.5, "opening_deg": 120, "candidate_facing_overlaps": list(Q3_OVERLAPS)},
        "length_scope": "north-south survey lines only; no turns or transits; no global optimality claim",
        "legacy_counterexample": old, "candidates": candidates, "selected_facing_overlap": best["facing_overlap"],
        "selected_line_count": best["line_count"], "selected_total_length_m": best["total_length_m"],
    }


def question4_audit(root: Path) -> dict:
    import openpyxl

    directory = root / "code" / "问题四"
    datasets = []
    source_depths = {}
    comparison = {"joined_coordinate_rows": 0, "Z_equals_negative_original_rows": 0, "Z1_equals_original_rows": 0, "Z1_equals_negative_original_rows": 0, "unmatched_coordinate_rows": 0}
    for filename in ("处理后数据1.xlsx", "处理后数据2.xlsx"):
        path = directory / filename
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = book.active
        values = sheet.iter_rows(values_only=True)
        header = next(values)
        minima, maxima = [None] * len(header), [None] * len(header)
        count, invalid, duplicates = 0, 0, 0
        coordinates = set()
        for row in values:
            count += 1
            if len(row) != len(header) or any(type(value) not in (int, float) or not isfinite(value) for value in row):
                invalid += 1
                continue
            for i, value in enumerate(row):
                minima[i] = value if minima[i] is None else min(minima[i], value)
                maxima[i] = value if maxima[i] is None else max(maxima[i], value)
            coordinate = row[0], row[1]
            duplicates += coordinate in coordinates
            coordinates.add(coordinate)
            if filename == "处理后数据1.xlsx":
                source_depths[coordinate] = row[2]
                continue
            if coordinate not in source_depths:
                comparison["unmatched_coordinate_rows"] += 1
                continue
            original = source_depths[coordinate]
            comparison["joined_coordinate_rows"] += 1
            comparison["Z_equals_negative_original_rows"] += row[2] == -original
            comparison["Z1_equals_original_rows"] += row[3] == original
            comparison["Z1_equals_negative_original_rows"] += row[3] == -original
        book.close()
        datasets.append({**file_record(path, root), "sheet": sheet.title, "header": header, "data_rows": count, "invalid_numeric_rows": invalid, "duplicate_coordinates": duplicates, "column_minima": minima, "column_maxima": maxima, "units": {"X": "nautical miles inferred from domain extent and source coordinates", "Y": "nautical miles inferred from domain extent and source coordinates", "Z": "metres; dataset 1 positive depth, dataset 2 negative height", "Z1": "dataset 2 only; generating rule and interpretation not supplied"}})
    return {
        "reconstruction_complete": False,
        "missing": ["original attachment workbook and its source identity", "generation rules for processed dataset 2", "traceable region fitting, clipping, and actual line lengths"],
        "datasets": datasets,
        "empirical_dataset_comparison": comparison,
        "source_scripts": [file_record(p, root) for p in sorted((root / "archive" / "competition-code" / "问题四").glob("*"))],
        "known_discrepancies": ["region 4 depth uses 6383.61 m but iteration stops at 6373.61 m", "region 1 length scaling uses 8869.43 m while spacing uses 8869.44 m", "region 3 length scaling uses 5599.81 m while spacing uses 5599.82 m", "region 4 length scaling uses both 5869.12 m and 5869.19 m", "legacy depth decrement uses sin(alpha) for a horizontal coordinate", "area-ratio length scaling does not clip individual survey lines to region boundaries"],
        "solver_status": "No Q4 solution produced. Existing workbooks are evidence, not reconstructed raw data.",
    }
